import argparse
import datetime as dt
import hashlib
import json
import re
import sys
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from pipeline import territories
from pipeline.common import (AGENCIES, AGENCY_FIELDS, AGENCY_ORDER, GAP_FIELDS, GAPS, LINE_FIELDS, LINES, SNAPSHOT,
                             load_regions, published_regions, read_json, today, write_csv, write_json)

SHEET = "ПЕРЕЧЕНЬ"
HEADER_ROW = 3
FIRST_ROW = 5
MIN_REGIONS, MAX_REGIONS, MIN_NUMBERS = 80, 95, 300
MAX_UNZIPPED = 100 * 1024 * 1024
SEPS = " -\u2013\u2014()"
PHONE_START = re.compile(r"(?<![\w+])(\+7|8)(?=[\s(\-\u2013\u2014]*\d)|(?<![\w(])\((?=\d{3,5}\))")
SHORT = re.compile(r"(?<![\w.:/\-])(1\d\d)(?![\w.:/])")
EXT = re.compile(r"\(?\s*((?:доб(?:авочный)?\.?[\s-]*\d[\d-]*)|(?:цифра\s*\d+))\s*\)?", re.I)
HOURS = re.compile(r"\(?\s*(?:[сc]\s+\d|[сc]\s+понедельник|в\s+рабоч|в\s+вечерн|в\s+выходн|в\s+праздн|в\s+будн|по\s+будн"
                   r"|в\s+режиме|\bрабочие\s+дни|(?:пн|вт|ср|чт|пт|сб|вс)\b|понедельник|режим\s+работы|время\s+работы|работает|круглосуточ|на\s+период"
                   r"|\d{1,2}:\d\d\s*-)", re.I)
TRIVIAL = re.compile(r"[\s,;()\-\u2013\u2014]*")
EDGE = "\t\n -\u2013\u2014:;,."
EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
FIO = re.compile(r"[А-ЯЁ][а-яё]+\s+[А-ЯЁ]\.\s?[А-ЯЁ]\.|[А-ЯЁ]\.\s?[А-ЯЁ]\.\s?[А-ЯЁ][а-яё]+"
                 r"|[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+(?:вич|вна|ична|чна)\b")
STAFF_TAIL = re.compile(r"ответственн\w*\s+сотрудник", re.I)
PERIOD = re.compile(r"на\s+период\s+с\s+\d{2}\.\d{2}\.\d{4}\s+по\s+(\d{2})\.(\d{2})\.(\d{4})")
TOPIC_WORDS = [
    ("staff", r"директор|заместител|\bзам\.|начальник|при[её]мн"),
    ("pain", r"обезбол|наркотическ"),
    ("pmp", r"паллиат"),
    ("drugs", r"лекарств|лекартсв|лек\.\s*обеспеч|препарат"),
    ("complaints", r"наруш|жалоб|обращени|прав\s+граждан|качеств"),
    ("general", r"горячая\s+линия|медицинской\s+помощ|мед\.\s*помощ|медико-санитарн|медицинских\s+услуг"),
    ("other", r"вакцин|коронавирус|covid|ковид|оплат\w*\s+труда|\bвмп\b|высокотехнолог|психолог|психиатр|полис|родовспом"),
]
TOLLFREE_CODES = ("800", "804")
TOPIC_ORDER = ["pain", "pmp", "drugs", "complaints", "general", "other", "staff"]


class FormatError(Exception):
    pass


@dataclass
class Item:
    raw: str
    tel: str
    line: int
    label: str
    heading: str
    head_line: int
    cont: list = field(default_factory=list)
    end_line: int = 0


@dataclass
class Stats:
    personal: int = 0
    email: int = 0
    notes: int = 0
    malformed: list = field(default_factory=list)
    merged: int = 0


def norm_space(s: str) -> str:
    return re.sub(r"\s+", " ", s.replace("\xa0", " ")).strip()


def take_phone(text: str, m: re.Match) -> tuple[int, str]:
    lead = m.group(0)
    digits = {"+7": "7", "8": "8"}.get(lead, "8")
    i = m.end() if lead in ("+7", "8") else m.start()
    end = i
    while i < len(text) and len(digits) < 11:
        ch = text[i]
        if ch.isdigit():
            digits += ch
            end = i + 1
        elif ch not in SEPS:
            break
        i += 1
    if len(digits) == 11 and not (end < len(text) and text[end].isdigit()):
        return end, "+7" + digits[1:]
    return end, ""


def find_numbers(line: str) -> list[dict]:
    spans, pos = [], 0
    for m in PHONE_START.finditer(line):
        if m.start() < pos:
            continue
        end, tel = take_phone(line, m)
        if tel or len(re.sub(r"\D", "", line[m.start():end])) >= 6:
            spans.append({"start": m.start(), "end": end, "raw": line[m.start():end].strip(), "tel": tel})
            pos = end
    masked = list(line)
    for s in spans:
        masked[s["start"]:s["end"]] = " " * (s["end"] - s["start"])
    masked = "".join(masked)
    for m in SHORT.finditer(masked):
        if re.search(r"доб\w*\.?\s*$", masked[max(0, m.start() - 14):m.start()], re.I):
            continue
        spans.append({"start": m.start(), "end": m.end(), "raw": m.group(1), "tel": m.group(1)})
    return sorted(spans, key=lambda s: s["start"])


def split_lines(cell: str) -> list[str]:
    out = []
    for ln in cell.replace("\xa0", " ").split("\n"):
        for part in re.split(r" {5,}", ln):
            out.append(norm_space(part))
    return out


def scrub(lines: list[str], stats: Stats) -> list[str]:
    out = []
    for i, ln in enumerate(lines):
        if STAFF_TAIL.search(ln):
            dropped = lines[i:]
            stats.personal += sum(len(find_numbers(x)) or 1 for x in dropped if x)
            break
        if FIO.search(ln):
            stats.personal += max(1, len(find_numbers(ln)))
            out.append("")
            continue
        if EMAIL.search(ln):
            stats.email += len(EMAIL.findall(ln))
            ln = norm_space(EMAIL.sub("", ln))
        out.append(ln)
    return out


def classify(line: str) -> str:
    if not line:
        return "blank"
    if find_numbers(line):
        return "number"
    if HOURS.match(line):
        return "hours"
    if line.endswith(":"):
        return "heading"
    return "text"


def next_kind(kinds: list[str], i: int) -> str:
    for k in kinds[i + 1:]:
        if k != "blank":
            return k
    return "end"


def number_items(line: str, idx: int, heading: str, head_line: int, stats: Stats, where: dict) -> list[Item]:
    spans = find_numbers(line)
    prefix = line[:spans[0]["start"]].strip(EDGE)
    tails = [line[s["end"]:(spans[k + 1]["start"] if k + 1 < len(spans) else len(line))] for k, s in enumerate(spans)]
    items, pending = [], []
    for k, s in enumerate(spans):
        pending.append(s)
        if k + 1 < len(spans) and TRIVIAL.fullmatch(tails[k]):
            continue
        label = " ".join(x for x in (prefix, tails[k].strip(EDGE)) if x)
        for p in pending:
            if not p["tel"]:
                stats.malformed.append({**where, "raw": p["raw"]})
                continue
            items.append(Item(raw=p["raw"], tel=p["tel"], line=idx, label=label, heading=heading,
                              head_line=head_line, end_line=idx))
        pending = []
    return items


def group_cell(cell: str, stats: Stats, where: dict) -> tuple[list[Item], list[str]]:
    lines = scrub(split_lines(cell), stats)
    kinds = [classify(ln) for ln in lines]
    items, group, prev = [], [], []
    heading, head_line = "", -1
    soft, soft_line = "", -1
    for i, (ln, kind) in enumerate(zip(lines, kinds)):
        if kind == "blank":
            heading, head_line, group, prev = "", -1, [], []
        elif kind == "heading" and next_kind(kinds, i) in ("number", "hours"):
            heading, head_line, group, prev = ln.rstrip(":").strip(), i, [], []
        elif kind == "hours" and (group if heading else prev):
            for it in group if heading else prev:
                it.cont.append(ln)
                it.end_line = i
        elif kind in ("text", "heading", "hours"):
            if next_kind(kinds, i) == "number" and kind != "hours":
                soft, soft_line = ln.rstrip(":").strip(), i
            else:
                stats.notes += 1
                prev = []
        else:
            head = " ".join(x for x in (heading, soft) if x)
            start = head_line if heading else (soft_line if soft else i)
            new = number_items(ln, i, head, start, stats, where)
            soft, soft_line = "", -1
            items += new
            prev = new
            if heading:
                group += new
    return items, lines


def split_label(label: str) -> tuple[str, str, str]:
    ext = ""
    m = EXT.search(label)
    if m:
        ext = m.group(1).strip()
        label = (label[:m.start()] + " " + label[m.end():]).strip()
    label = norm_space(label).strip(EDGE)
    if label.startswith("(") and label.endswith(")") and label.count("(") == 1:
        label = label[1:-1].strip(EDGE)
    m = HOURS.search(label)
    if not m:
        return label.strip(EDGE + "("), "", ext
    topic = label[:m.start()].strip(EDGE + "(")
    hours = label[m.start():].strip(EDGE)
    if hours.startswith("(") and hours.endswith(")"):
        hours = hours[1:-1].strip(EDGE)
    elif hours.startswith("(") and ")" not in hours:
        hours = hours[1:].strip(EDGE)
    elif hours.endswith(")") and hours.count("(") < hours.count(")"):
        hours = hours[:-1].strip(EDGE)
    return topic, hours, ext


def is_24x7(text: str) -> bool:
    for m in re.finditer(r"круглосуточ\w*", text, re.I):
        if not re.match(r"\s+в\s+(выходн|праздн)", text[m.end():], re.I):
            return True
    return False


def topics_of(text: str, hours: str, assured: str) -> list[str]:
    low = text.lower()
    found = {name for name, rx in TOPIC_WORDS if re.search(rx, low)}
    m = PERIOD.search(hours)
    if m and dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1))).isoformat() < assured:
        return ["other"]
    if "staff" in found:
        return ["staff"]
    if found & {"pain", "pmp"}:
        found.discard("general")
    if not found:
        found = {"general"}
    return [t for t in TOPIC_ORDER if t in found]


def audience_of(text: str) -> str:
    low = text.lower()
    adults, children = "взросл" in low, re.search(r"\bдет|детск", low) is not None
    return "adults" if adults and not children else "children" if children and not adults else "all"


def kind_of(tel: str) -> str:
    if len(tel) == 3:
        return "short"
    if tel[2:5] in TOLLFREE_CODES:
        return "tollfree"
    return "mobile" if tel[2] == "9" else "landline"


def display(raw: str, tel: str) -> str:
    if len(tel) == 3:
        return tel
    body = tel[2:]
    m = re.search(r"\((\d{3,5})\)", raw)
    code_len = len(m.group(1)) if m else len(re.findall(r"\d+", re.sub(r"^\s*(\+7|8)", "", raw))[0])
    if body[:3] == "800" or body[0] == "9" or not 3 <= code_len <= 5:
        code_len = 3
    code, rest = body[:code_len], body[code_len:]
    groups = {7: (3, 2, 2), 6: (2, 2, 2), 5: (1, 2, 2)}[len(rest)]
    parts, pos = [], 0
    for g in groups:
        parts.append(rest[pos:pos + g])
        pos += g
    return f"8 ({code}) {'-'.join(parts)}"


def agency_of(name: str) -> str:
    if "росздравнадзор" in name.lower():
        return "rzn"
    if re.search(r"фонд\w*.*страхован|\bОМС\b|медицинское\s+страхование", name, re.I):
        return "tfoms"
    return "minzdrav"


def finalize(it: Item, lines: list[str], ctx: dict) -> dict:
    topic, hours, ext = split_label(it.label)
    if it.heading:
        topic = ", ".join(x for x in (it.heading.strip(EDGE), topic) if x)
    hours = "; ".join(x for x in [hours] + [c.strip(EDGE) for c in it.cont] if x)
    blob = " ".join([it.heading, it.label] + it.cont)
    source_text = " ".join(x for x in lines[it.head_line:it.end_line + 1] if x)
    topics = topics_of(topic, hours, ctx["assured_date"])
    return {
        "region_id": ctx["region_id"], "agency": ctx["agency"], "agency_name": ctx["agency_name"],
        "number_raw": it.raw, "number_display": display(it.raw, it.tel), "tel": it.tel, "ext_text": ext,
        "kind": kind_of(it.tel), "topic_text": topic, "hours_text": hours, "is_24x7": is_24x7(blob),
        "topics": topics, "audience": audience_of(topic), "source_row": ctx["row"],
        "source_text": source_text, "assured_date": ctx["assured_date"],
    }


def dedupe(rows: list[dict], stats: Stats) -> list[dict]:
    out = []
    for r in rows:
        twin = next((o for o in out if o["tel"] == r["tel"] and o["topic_text"] == r["topic_text"]), None)
        if twin is None:
            out.append(r)
            continue
        stats.merged += 1
        if len(r["hours_text"]) + r["is_24x7"] > len(twin["hours_text"]) + twin["is_24x7"]:
            out[out.index(twin)] = r
    return out


def check_zip(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as z:
            total = sum(i.file_size for i in z.infolist())
    except (zipfile.BadZipFile, OSError) as e:
        raise FormatError(f"Файл {path.name} не читается как xlsx: {e}") from e
    if total > MAX_UNZIPPED:
        raise FormatError(f"Файл {path.name} после распаковки займёт {total} байт, предел {MAX_UNZIPPED}. Разбор остановлен")


def load_sheet(path: Path):
    import openpyxl
    from openpyxl import xml as oxml
    if not oxml.DEFUSEDXML:
        raise FormatError("Нет пакета defusedxml: разбор xlsx без защиты XML не запускаем")
    check_zip(path)
    try:
        wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    except Exception as e:
        raise FormatError(f"Файл {path.name} не читается как xlsx: {e}") from e
    if SHEET not in wb.sheetnames:
        raise FormatError(f"В таблице нет листа «{SHEET}», есть только: {', '.join(wb.sheetnames)}")
    ws = wb[SHEET]
    head = [str(ws.cell(HEADER_ROW, c).value or "").strip() for c in range(1, max(ws.max_column, 4) + 1)]
    if not head[1].startswith("Регион") or "Контактные номера" not in head[3]:
        raise FormatError(f"Строка {HEADER_ROW}: ожидались столбцы «Регионы», «Полное наименование», «Контактные номера»")
    extra = [i + 1 for i, v in enumerate(head) if i >= 4 and v]
    if extra:
        raise FormatError(f"В таблице появились лишние столбцы: {extra}. Разбор остановлен")
    return ws


def assured_from(path: Path, fallback: str) -> str:
    m = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", path.name) or None
    if m:
        return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", path.name)
    return m.group(0) if m else fallback


def iter_rows(ws):
    region = ""
    for r in range(FIRST_ROW, ws.max_row + 1):
        b, c, d = (ws.cell(r, k).value for k in (2, 3, 4))
        if b is not None and str(b).strip():
            region = str(b).strip()
        if c is None and d is None:
            continue
        if not region:
            raise FormatError(f"Строка {r}: не указан регион, а ячейки региона должны быть объединены как раньше")
        yield r, region, str(c or "").strip(), "" if d is None else str(d)


def parse(path: Path, regions: list[dict], assured_date: str) -> dict:
    ws = load_sheet(path)
    by_name = {r["name_source"]: r for r in regions}
    stats = Stats()
    lines, names, cells, agencies, gaps = [], [], [], [], []
    counts = {"rows": 0, "rzn": 0, "minzdrav": 0, "tfoms": 0, "excluded_rows": 0, "excluded_numbers": 0}
    excluded_names: set[str] = set()
    for row, name, agency_name, cell in iter_rows(ws):
        if name not in names:
            names.append(name)
        agency = agency_of(agency_name)
        counts["rows"] += 1
        counts[agency] += 1
        cells.append(cell)
        if territories.is_excluded(name, agency_name):
            counts["excluded_rows"] += 1
            counts["excluded_numbers"] += sum(len(find_numbers(x)) for x in split_lines(cell))
            excluded_names.add(name)
            continue
        if name not in by_name:
            raise FormatError(f"Строка {row}: регион «{name}» не найден в data/regions.csv (поле name_source)")
        reg = by_name[name]
        agencies.append({"region_id": reg["id"], "agency": agency, "agency_name": norm_space(agency_name),
                         "source_row": row})
        where = {"region_id": reg["id"], "agency": agency, "source_row": row}
        before = len(stats.malformed)
        items, clean = group_cell(cell, stats, where)
        for m in stats.malformed[before:]:
            gaps.append({**where, "kind": "malformed", "text_verbatim": m["raw"]})
        if not items and before == len(stats.malformed):
            gaps.append({**where, "kind": "no_number", "text_verbatim": norm_space(" ".join(x for x in clean if x))})
        ctx = {"region_id": reg["id"], "agency": agency, "agency_name": norm_space(agency_name), "row": row,
               "assured_date": assured_date}
        lines += dedupe([finalize(it, clean, ctx) for it in items], stats)
    if not MIN_REGIONS <= len(names) <= MAX_REGIONS:
        raise FormatError(f"Регионов в таблице {len(names)}, ожидалось от {MIN_REGIONS} до {MAX_REGIONS}")
    if len(lines) < MIN_NUMBERS:
        raise FormatError(f"Номеров разобрано {len(lines)}, ожидалось не меньше {MIN_NUMBERS}")
    assign_ids(lines, regions)
    agencies.sort(key=lambda a: (a["region_id"], AGENCY_ORDER[a["agency"]]))
    counts["excluded_regions"] = len(excluded_names)
    gaps.sort(key=lambda g: (g["region_id"], AGENCY_ORDER[g["agency"]], g["source_row"]))
    return {"lines": lines, "agencies": agencies, "gaps": gaps, "regions_total": len(names), "counts": counts, "stats": stats,
            "cells": {"pall": sum("паллиат" in c.lower() for c in cells),
                      "round_clock": sum("круглосуточ" in c.lower() for c in cells),
                      "n122": sum(bool(re.search(r"(?<!\d)122(?!\d)", c)) for c in cells)}}


def assign_ids(lines: list[dict], regions: list[dict]) -> None:
    slug = {r["id"]: r["slug"] for r in regions if r["id"]}
    owners: dict[str, set] = {}
    for ln in lines:
        if len(ln["tel"]) > 3:
            owners.setdefault(ln["tel"], set()).add(ln["region_id"])
    lines.sort(key=lambda x: (x["region_id"], AGENCY_ORDER[x["agency"]], x["source_row"]))
    seen: dict[str, int] = {}
    for ln in lines:
        base = f"{slug[ln['region_id']]}-{ln['agency']}-{'8' + ln['tel'][2:] if len(ln['tel']) > 3 else ln['tel']}"
        seen[base] = seen.get(base, 0) + 1
        ln["id"] = base if seen[base] == 1 else f"{base}-{seen[base]}"
        ln["shared_with"] = sorted(owners.get(ln["tel"], set()) - {ln["region_id"]})


def to_csv_rows(lines: list[dict]) -> list[dict]:
    out = []
    for ln in lines:
        row = dict(ln)
        row["is_24x7"] = "true" if ln["is_24x7"] else "false"
        row["topics"] = ";".join(ln["topics"])
        row["shared_with"] = ";".join(ln["shared_with"])
        out.append(row)
    return out


def region_hashes(lines: list[dict]) -> dict:
    out: dict[str, list] = {}
    for ln in lines:
        out.setdefault(ln["region_id"], []).append({k: ln[k] for k in LINE_FIELDS})
    return {k: hashlib.sha256(json.dumps(v, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
            for k, v in sorted(out.items())}


def snapshot_counts(res: dict, regions: list[dict]) -> dict:
    st = res["stats"]
    return {
        "rows": res["counts"]["rows"],
        "rows_by_agency": {k: res["counts"][k] for k in ("rzn", "minzdrav", "tfoms")},
        "regions_total": res["regions_total"],
        "regions_published": len(published_regions(regions)),
        "lines": len(res["lines"]),
        "cells": res["cells"],
        "dropped": {"personal": st.personal, "email": st.email, "notes": st.notes, "merged_duplicates": st.merged,
                    "excluded_regions": res["counts"]["excluded_regions"],
                    "excluded_rows": res["counts"]["excluded_rows"],
                    "excluded_numbers": res["counts"]["excluded_numbers"]},
        "malformed": st.malformed,
        "region_hashes": region_hashes(res["lines"]),
    }


def run(root: Path, xlsx: Path, *, source: dict | None = None) -> dict:
    regions = load_regions(root)
    snap = read_json(root / SNAPSHOT, {}) or {}
    sha = hashlib.sha256(xlsx.read_bytes()).hexdigest()
    src = dict(source or {})
    if sha != snap.get("sha256") and not source:
        src = {"sha256": sha, "bytes": xlsx.stat().st_size}
    assured = assured_from(Path(src.get("url", "")) if src.get("url") else xlsx, snap.get("assured_date") or today())
    res = parse(xlsx, regions, assured)
    write_csv(root / LINES, LINE_FIELDS, to_csv_rows(res["lines"]))
    write_csv(root / AGENCIES, AGENCY_FIELDS, res["agencies"])
    write_csv(root / GAPS, GAP_FIELDS, res["gaps"])
    new = {**snap, **src, "sha256": sha, "bytes": xlsx.stat().st_size, "assured_date": assured,
           **snapshot_counts(res, regions)}
    write_json(root / SNAPSHOT, new)
    return new


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="python -m pipeline.parse_rzn")
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--url")
    ap.add_argument("--etag")
    ap.add_argument("--last-modified")
    a = ap.parse_args(argv)
    source = {k: v for k, v in (("url", a.url), ("etag", a.etag), ("last_modified", a.last_modified)) if v}
    try:
        snap = run(Path.cwd(), a.xlsx, source=source or None)
    except FormatError as e:
        print(f"Ошибка формата: {e}", file=sys.stderr)
        return 1
    print(f"Разобрано номеров: {snap['lines']}, регионов в таблице: {snap['regions_total']}, "
          f"опубликуем: {snap['regions_published']}, записано в {LINES}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
