import csv
import re
import shutil
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
CACHE = ROOT / "raw" / "cache"


STREET = re.compile(r"(?:ул|пл|пр|просп|пер|наб|б-р|бульв|имени|им)\.?\s*$", re.I)
RZN_EMAIL = re.compile(r"[\w.+-]+@(?:[\w-]+\.)*roszdravnadzor\.gov\.ru")


def personal(text: str) -> list[str]:
    from pipeline.parse_rzn import FIO
    return [m.group(0) for m in FIO.finditer(text) if not STREET.search(text[max(0, m.start() - 8):m.start()])]


def strings(obj):
    if isinstance(obj, str):
        yield obj
    elif isinstance(obj, dict):
        for v in obj.values():
            yield from strings(v)
    elif isinstance(obj, list):
        for v in obj:
            yield from strings(v)


def real_xlsx() -> Path | None:
    found = sorted(p for p in CACHE.glob("*.xlsx") if "03.03.2026" in p.name or "2026-03-03" in p.name)
    return found[0] if found else None


def read_rows(path: Path) -> list[dict]:
    with open(path, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def synth_cells(root: Path = ROOT) -> list[tuple[str, str, str]]:
    regions = read_rows(root / "data" / "regions.csv")
    lines = read_rows(root / "data" / "rzn" / "lines.csv")
    agencies = read_rows(root / "data" / "rzn" / "agencies.csv")
    out = []
    for r in regions:
        for a in (x for x in agencies if x["region_id"] == r["id"]):
            own = [x for x in lines if x["region_id"] == r["id"] and x["agency"] == a["agency"]]
            cell = "\n".join(x["number_raw"] + (f" - {x['topic_text']}" if x["topic_text"] else "")
                             + (f" ({x['hours_text']})" if x["hours_text"] else "") for x in own)
            out.append((r["name_source"], a["agency_name"], cell or "нет номера"))
    return out


def write_xlsx(path: Path, cells, *, sheet="ПЕРЕЧЕНЬ", extra_col=False, drop_first_region=False) -> Path:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.cell(2, 2, "Номера телефонов «Горячих линий» в субъектах Российской Федерации")
    ws.cell(3, 2, "Регионы ")
    ws.cell(3, 3, "Полное наименование Территориального органа Росздравнадзора")
    ws.cell(3, 4, "Контактные номера")
    if extra_col:
        ws.cell(3, 5, "Примечание")
    row, prev = 5, None
    for name, agency, cell in cells:
        if name != prev and not (drop_first_region and row == 5):
            ws.cell(row, 2, name)
        prev = name
        ws.cell(row, 3, agency)
        ws.cell(row, 4, cell)
        row += 1
    wb.save(path)
    return path


@pytest.fixture
def tmp_root(tmp_path):
    shutil.copytree(ROOT / "data", tmp_path / "data")
    return tmp_path
