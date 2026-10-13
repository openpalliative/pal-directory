import re

import pytest
import yaml

from pipeline import parse_rzn
from pipeline.common import load_regions
from pipeline.tests.conftest import ROOT, read_rows, real_xlsx

XLSX = real_xlsx()
pytestmark = pytest.mark.skipif(XLSX is None, reason="нет raw/cache/горячая линия-03.03.2026.xlsx, golden пропущен")
GOLD = yaml.safe_load((ROOT / "tests" / "golden" / "rzn-counts.yaml").read_text(encoding="utf-8"))
SPOT = yaml.safe_load((ROOT / "tests" / "golden" / "spot-lines.yaml").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def res():
    return parse_rzn.parse(XLSX, load_regions(ROOT), "2026-03-03")


@pytest.fixture(scope="module")
def cells():
    import openpyxl
    ws = openpyxl.load_workbook(XLSX)["ПЕРЕЧЕНЬ"]
    return {r: parse_rzn.norm_space(str(ws.cell(r, 4).value or "")) for r in range(5, ws.max_row + 1)}


def test_counts(res):
    c = res["counts"]
    assert res["regions_total"] == GOLD["regions_total"]
    assert c["rows"] == GOLD["rows"]
    assert {k: c[k] for k in ("rzn", "minzdrav", "tfoms")} == GOLD["rows_by_agency"]
    assert c["excluded_rows"] == GOLD["excluded_rows"]
    assert len(res["lines"]) == GOLD["lines"]
    assert res["cells"] == {"pall": GOLD["cells_pall"], "round_clock": GOLD["cells_round_clock"], "n122": GOLD["cells_122"]}
    assert len(res["stats"].malformed) == GOLD["malformed"]
    assert res["stats"].email == GOLD["email_dropped"]
    assert res["stats"].personal >= GOLD["personal_dropped_min"]


def test_published_regions(res):
    regions = load_regions(ROOT)
    published = {r["id"] for r in regions}
    assert len(published) == GOLD["regions_published"]
    assert res["counts"]["excluded_regions"] == GOLD["regions_excluded"]
    assert {ln["region_id"] for ln in res["lines"]} == published


def test_committed_csv_matches_parse(res):
    committed = read_rows(ROOT / "data" / "rzn" / "lines.csv")
    fresh = parse_rzn.to_csv_rows(res["lines"])
    assert [r["id"] for r in committed] == [r["id"] for r in fresh]
    for a, b in zip(committed, fresh):
        assert a == {k: str(v) for k, v in b.items()}, a["id"]


def test_invariants(res, cells):
    for ln in res["lines"]:
        assert re.fullmatch(r"\+7\d{10}|\d{3}", ln["tel"]), ln["id"]
        assert ln["source_text"] in cells[ln["source_row"]], ln["id"]
        assert ln["number_raw"] in ln["source_text"], ln["id"]
        digits = re.sub(r"\D", "", ln["number_raw"])
        assert digits[-10:] == ln["tel"][-10:] or ln["tel"] == digits, ln["id"]
        if ln["is_24x7"]:
            assert "круглосуточ" in ln["source_text"].lower(), ln["id"]


def test_spot_lines(res):
    by_id = {ln["id"]: ln for ln in res["lines"]}
    for spot in SPOT["lines"]:
        ln = by_id.get(spot["id"])
        assert ln is not None, spot["id"]
        for key, want in spot["expect"].items():
            if key == "topics_has":
                assert want in ln["topics"], spot["id"]
            else:
                assert ln[key] == want, f"{spot['id']}.{key}: {ln[key]!r} != {want!r}"
    for spot in SPOT["absent"]:
        assert spot not in by_id


def test_region_rules(res):
    vgg = {ln["tel"]: ln["is_24x7"] for ln in res["lines"] if ln["region_id"] == "RU-VGG"}
    assert vgg["+78442951513"] is False and vgg["+78442951520"] is False and vgg["+78442598575"] is True
    drugs = next(ln for ln in res["lines"] if ln["tel"] == "+78442951513")
    assert drugs["hours_text"].endswith("пт. с 16-00 до 08-00, круглосуточно в выходные и праздничные дни")
    sar = [ln for ln in res["lines"] if ln["tel"] == "+78452799312"]
    assert len(sar) == 1 and sar[0]["is_24x7"] is True
    yan = [ln for ln in res["lines"] if ln["tel"] == "+73492237640"]
    assert len(yan) == 2 and yan[1]["topics"] == ["other"]
